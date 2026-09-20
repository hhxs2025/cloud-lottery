# -*- coding: utf-8 -*-
import sys
import os

vendor_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'vendor')
if os.path.exists(vendor_dir):
    sys.path.insert(0, vendor_dir)

import sqlite3
import random
import uuid
from datetime import datetime
from functools import wraps

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from flask import Flask, render_template, request, redirect, url_for, flash, session, jsonify, send_from_directory
from werkzeug.serving import make_server
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.secret_key = 'ynos-cloud-lottery-secret-key-2026'

# ---------- 路径 ----------
PKG_VAR = os.environ.get('TRIM_PKGVAR')
if PKG_VAR:
    DATA_DIR = os.path.join(PKG_VAR, 'data')
    UPLOAD_DIR = os.path.join(PKG_VAR, 'uploads')
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    DATA_DIR = os.path.join(BASE_DIR, 'data')
    UPLOAD_DIR = os.path.join(BASE_DIR, 'static', 'uploads')

AVATAR_DIR = os.path.join(UPLOAD_DIR, 'avatars')
DB_PATH = os.path.join(DATA_DIR, 'lottery.db')

ALLOWED_IMAGE_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'webp'}
ADMIN_USERNAME = os.environ.get('LOTTERY_ADMIN_USERNAME') or 'admin'
ADMIN_PASSWORD = os.environ.get('LOTTERY_ADMIN_PASSWORD') or 'admin123'

for d in [DATA_DIR, UPLOAD_DIR, AVATAR_DIR]:
    os.makedirs(d, exist_ok=True)

# ---------- 数据库 ----------
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    c = conn.cursor()
    c.executescript('''
        CREATE TABLE IF NOT EXISTS config (key TEXT PRIMARY KEY, value TEXT);
        CREATE TABLE IF NOT EXISTS candidates (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL UNIQUE, photo_path TEXT);
        CREATE TABLE IF NOT EXISTS prizes (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, prize_desc TEXT, quantity INTEGER NOT NULL DEFAULT 1, sort_order INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS predetermined (id INTEGER PRIMARY KEY AUTOINCREMENT, prize_id INTEGER NOT NULL, candidate_id INTEGER NOT NULL, UNIQUE(prize_id, candidate_id));
        CREATE TABLE IF NOT EXISTS winners (id INTEGER PRIMARY KEY AUTOINCREMENT, prize_id INTEGER NOT NULL, candidate_id INTEGER NOT NULL, draw_order INTEGER NOT NULL, created_at TEXT NOT NULL);
    ''')
    c.execute("INSERT OR IGNORE INTO config (key, value) VALUES ('title', '年会抽奖')")
    c.execute("INSERT OR IGNORE INTO config (key, value) VALUES ('title_color', '#222')")
    c.execute("INSERT OR IGNORE INTO config (key, value) VALUES ('title_size', '48')")
    c.execute("INSERT OR IGNORE INTO config (key, value) VALUES ('title_font', 'default')")
    c.execute("INSERT OR IGNORE INTO config (key, value) VALUES ('subtitle', '幸运大抽奖')")
    c.execute("INSERT OR IGNORE INTO config (key, value) VALUES ('subtitle_color', '#666')")
    c.execute("INSERT OR IGNORE INTO config (key, value) VALUES ('subtitle_size', '24')")
    conn.commit()
    conn.close()
init_db()

# ---------- 工具 ----------
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get('admin_logged_in'):
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

def allowed_image(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_IMAGE_EXTENSIONS

def get_config(key, default=''):
    conn = get_db()
    row = conn.execute("SELECT value FROM config WHERE key=?", (key,)).fetchone()
    conn.close()
    return row['value'] if row else default

def set_config(key, value):
    conn = get_db()
    conn.execute("INSERT OR REPLACE INTO config (key, value) VALUES (?, ?)", (key, value))
    conn.commit()
    conn.close()

def get_all_candidates():
    conn = get_db()
    rows = conn.execute("SELECT * FROM candidates ORDER BY id").fetchall()
    conn.close()
    return [dict(r) for r in rows]

def get_all_prizes():
    conn = get_db()
    rows = conn.execute("SELECT * FROM prizes ORDER BY sort_order, id").fetchall()
    conn.close()
    return [dict(r) for r in rows]

def get_prize(pid):
    conn = get_db()
    row = conn.execute("SELECT * FROM prizes WHERE id=?", (pid,)).fetchone()
    conn.close()
    return dict(row) if row else None

def get_winners_for_prize(pid):
    conn = get_db()
    rows = conn.execute('''
        SELECT w.*, c.name, c.photo_path
        FROM winners w
        JOIN candidates c ON w.candidate_id = c.id
        WHERE w.prize_id=?
        ORDER BY w.draw_order
    ''', (pid,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

def get_predetermined_for_prize(pid):
    conn = get_db()
    rows = conn.execute('''
        SELECT p.*, c.name, c.photo_path
        FROM predetermined p
        JOIN candidates c ON p.candidate_id = c.id
        WHERE p.prize_id=?
    ''', (pid,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

def get_all_winners():
    conn = get_db()
    rows = conn.execute('''
        SELECT w.*, c.name, c.photo_path, pr.name as prize_name, pr.prize_desc
        FROM winners w
        JOIN candidates c ON w.candidate_id = c.id
        JOIN prizes pr ON w.prize_id = pr.id
        ORDER BY w.created_at, w.draw_order
    ''').fetchall()
    conn.close()
    return [dict(r) for r in rows]

# ---------- 路由 ----------
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        u = request.form.get('username', '').strip()
        p = request.form.get('password', '').strip()
        if u == ADMIN_USERNAME and p == ADMIN_PASSWORD:
            session['admin_logged_in'] = True
            return redirect(url_for('admin'))
        flash('用户名或密码错误', 'error')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('index'))

@app.route('/')
def index():
    return render_template('index.html',
        title=get_config('title', '年会抽奖'),
        title_color=get_config('title_color', '#222'),
        title_size=get_config('title_size', '48'),
        title_font=get_config('title_font', 'default'),
        subtitle=get_config('subtitle', '幸运大抽奖'),
        subtitle_color=get_config('subtitle_color', '#666'),
        subtitle_size=get_config('subtitle_size', '24'),
        prizes=get_all_prizes(),
        candidates=get_all_candidates(),
        has_data=bool(get_all_prizes() and get_all_candidates())
    )

@app.route('/admin')
@login_required
def admin():
    prizes = get_all_prizes()
    candidates = get_all_candidates()
    winners = get_all_winners()
    for p in prizes:
        p['predetermined'] = get_predetermined_for_prize(p['id'])
        p['winners'] = get_winners_for_prize(p['id'])
        p['remaining'] = max(0, p['quantity'] - len(p['winners']))
    return render_template('admin.html',
        title=get_config('title', '年会抽奖'),
        title_color=get_config('title_color', '#222'),
        title_size=get_config('title_size', '48'),
        title_font=get_config('title_font', 'default'),
        subtitle=get_config('subtitle', '幸运大抽奖'),
        subtitle_color=get_config('subtitle_color', '#666'),
        subtitle_size=get_config('subtitle_size', '24'),
        prizes=prizes,
        candidates=candidates,
        winners=winners
    )

# ---------- API ----------
@app.route('/api/config', methods=['POST'])
@login_required
def save_config():
    keys = ['title', 'title_color', 'title_size', 'title_font', 'subtitle', 'subtitle_color', 'subtitle_size']
    for key in keys:
        value = request.form.get(key, '').strip()
        if value:
            set_config(key, value)
    flash('设置已保存', 'success')
    return redirect(url_for('admin'))

@app.route('/api/import_candidates', methods=['POST'])
@login_required
def import_candidates():
    action = request.form.get('action', 'append')
    conn = get_db()
    try:
        if action == 'replace':
            conn.execute("DELETE FROM candidates")
            conn.execute("DELETE FROM predetermined")
            conn.execute("DELETE FROM winners")
        if 'file' in request.files:
            file = request.files['file']
            if file and file.filename:
                ext = file.filename.rsplit('.', 1)[-1].lower()
                names = []
                if ext == 'txt':
                    content = file.read().decode('utf-8-sig')
                    names = [line.strip() for line in content.splitlines() if line.strip()]
                elif ext in ('xlsx', 'xls'):
                    if load_workbook is None:
                        flash('openpyxl 未安装，无法解析 Excel 文件', 'error')
                        return redirect(url_for('admin'))
                    file.seek(0)
                    wb = load_workbook(file)
                    ws = wb.active
                    for row in ws.iter_rows(min_row=1, values_only=True):
                        v = str(row[0]) if row and row[0] is not None else ''
                        v = v.strip()
                        if v and v.lower() != 'nan':
                            names.append(v)
                else:
                    flash('不支持的文件格式', 'error')
                    return redirect(url_for('admin'))
                for name in names:
                    try:
                        conn.execute("INSERT INTO candidates (name) VALUES (?)", (name,))
                    except sqlite3.IntegrityError:
                        pass
                conn.commit()
                flash('名单导入成功', 'success')
        else:
            flash('请上传文件', 'error')
    except Exception as e:
        conn.rollback()
        flash(f'导入失败：{str(e)}', 'error')
    finally:
        conn.close()
    return redirect(url_for('admin'))

@app.route('/api/candidate/<int:cid>/photo', methods=['POST'])
@login_required
def upload_candidate_photo(cid):
    if 'photo' not in request.files:
        flash('未选择文件', 'error')
        return redirect(url_for('admin'))
    file = request.files['photo']
    if not file or not file.filename:
        flash('未选择文件', 'error')
        return redirect(url_for('admin'))
    if not allowed_image(file.filename):
        flash('仅支持图片文件', 'error')
        return redirect(url_for('admin'))

    conn = get_db()
    row = conn.execute("SELECT photo_path FROM candidates WHERE id=?", (cid,)).fetchone()
    if row and row['photo_path']:
        old_full = os.path.join(UPLOAD_DIR, os.path.basename(row['photo_path']))
        try:
            if os.path.exists(old_full):
                os.remove(old_full)
        except Exception:
            pass

    ext = file.filename.rsplit('.', 1)[1].lower()
    filename = f"avatar_{cid}_{uuid.uuid4().hex[:8]}.{ext}"
    save_path = os.path.join(AVATAR_DIR, filename)
    file.save(save_path)

    conn.execute("UPDATE candidates SET photo_path=? WHERE id=?",
                 (f"uploads/avatars/{filename}", cid))
    conn.commit()
    conn.close()
    flash('照片已上传', 'success')
    return redirect(url_for('admin'))

@app.route('/api/candidate/<int:cid>', methods=['DELETE'])
@login_required
def delete_candidate(cid):
    conn = get_db()
    row = conn.execute("SELECT photo_path FROM candidates WHERE id=?", (cid,)).fetchone()
    if row and row['photo_path']:
        old_full = os.path.join(UPLOAD_DIR, os.path.basename(row['photo_path']))
        try:
            if os.path.exists(old_full):
                os.remove(old_full)
        except Exception:
            pass
    conn.execute("DELETE FROM candidates WHERE id=?", (cid,))
    conn.execute("DELETE FROM predetermined WHERE candidate_id=?", (cid,))
    conn.execute("DELETE FROM winners WHERE candidate_id=?", (cid,))
    conn.commit()
    conn.close()
    return jsonify({'success': True})

@app.route('/api/prize', methods=['POST'])
@login_required
def save_prize():
    data = request.get_json() or request.form
    name = (data.get('name') or '').strip()
    desc = (data.get('prize_desc') or '').strip()
    try:
        qty = int(data.get('quantity', 1))
    except ValueError:
        qty = 1
    if not name:
        return jsonify({'success': False, 'message': '奖项名称不能为空'})
    conn = get_db()
    conn.execute("INSERT INTO prizes (name, prize_desc, quantity) VALUES (?, ?, ?)", (name, desc, qty))
    conn.commit()
    conn.close()
    return jsonify({'success': True})

@app.route('/api/prize/<int:pid>', methods=['PUT', 'DELETE'])
@login_required
def prize_action(pid):
    conn = get_db()
    if request.method == 'DELETE':
        conn.execute("DELETE FROM winners WHERE prize_id=?", (pid,))
        conn.execute("DELETE FROM predetermined WHERE prize_id=?", (pid,))
        conn.execute("DELETE FROM prizes WHERE id=?", (pid,))
        conn.commit()
        conn.close()
        return jsonify({'success': True})
    data = request.get_json() or request.form
    name = (data.get('name') or '').strip()
    desc = (data.get('prize_desc') or '').strip()
    try:
        qty = int(data.get('quantity', 1))
    except ValueError:
        qty = 1
    if not name:
        return jsonify({'success': False, 'message': '奖项名称不能为空'})
    existing = conn.execute("SELECT COUNT(*) as cnt FROM predetermined WHERE prize_id=?", (pid,)).fetchone()['cnt']
    if qty < existing:
        conn.close()
        return jsonify({'success': False, 'message': f'名额不能小于当前内定人数（{existing}）'})
    conn.execute("UPDATE prizes SET name=?, prize_desc=?, quantity=? WHERE id=?", (name, desc, qty, pid))
    conn.commit()
    conn.close()
    return jsonify({'success': True})

@app.route('/api/prize/<int:pid>/predetermined', methods=['POST', 'DELETE'])
@login_required
def predetermined_action(pid):
    data = request.get_json() or request.form
    cid = data.get('candidate_id')
    try:
        cid = int(cid)
    except (ValueError, TypeError):
        return jsonify({'success': False, 'message': '参数错误'})
    conn = get_db()
    prize = conn.execute("SELECT * FROM prizes WHERE id=?", (pid,)).fetchone()
    if not prize:
        conn.close()
        return jsonify({'success': False, 'message': '奖项不存在'})
    existing = conn.execute("SELECT COUNT(*) as cnt FROM predetermined WHERE prize_id=?", (pid,)).fetchone()['cnt']
    if request.method == 'POST':
        if existing >= prize['quantity']:
            conn.close()
            return jsonify({'success': False, 'message': '内定人数不能超过该奖项名额'})
        try:
            conn.execute("INSERT INTO predetermined (prize_id, candidate_id) VALUES (?, ?)", (pid, cid))
            conn.commit()
        except sqlite3.IntegrityError:
            conn.close()
            return jsonify({'success': False, 'message': '该人员已是本奖项内定人'})
    else:
        conn.execute("DELETE FROM predetermined WHERE prize_id=? AND candidate_id=?", (pid, cid))
        conn.commit()
    conn.close()
    return jsonify({'success': True})

@app.route('/api/draw/<int:pid>', methods=['POST'])
def draw_prize(pid):
    prize = get_prize(pid)
    if not prize:
        return jsonify({'success': False, 'message': '奖项不存在'})
    candidates = get_all_candidates()
    if not candidates:
        return jsonify({'success': False, 'message': '请登录后台导入名单并设定奖项信息。'})
    conn = get_db()
    all_winner_ids = [r['candidate_id'] for r in conn.execute("SELECT candidate_id FROM winners").fetchall()]
    current_winners = get_winners_for_prize(pid)
    if len(current_winners) >= prize['quantity']:
        conn.close()
        return jsonify({'success': False, 'message': '该奖项已抽满'})
    predetermined = get_predetermined_for_prize(pid)
    predetermined_ids = {p['candidate_id'] for p in predetermined}
    predetermined_candidates = [c for c in candidates if c['id'] in predetermined_ids and c['id'] not in all_winner_ids]
    chosen = None
    if predetermined_candidates:
        chosen = random.choice(predetermined_candidates)
    else:
        eligible = [c for c in candidates if c['id'] not in all_winner_ids]
        if not eligible:
            conn.close()
            return jsonify({'success': False, 'message': '没有可抽奖人员（所有人均已中奖）'})
        chosen = random.choice(eligible)
    draw_order = len(current_winners) + 1
    now = datetime.now().isoformat()
    conn.execute("INSERT INTO winners (prize_id, candidate_id, draw_order, created_at) VALUES (?, ?, ?, ?)",
                 (pid, chosen['id'], draw_order, now))
    conn.commit()
    conn.close()
    return jsonify({
        'success': True,
        'candidate': {
            'id': chosen['id'],
            'name': chosen['name'],
            'photo_path': chosen.get('photo_path', '')
        },
        'draw_order': draw_order
    })

@app.route('/api/state')
def state():
    prizes = get_all_prizes()
    candidates = get_all_candidates()
    winners = get_all_winners()
    for p in prizes:
        p['winners'] = [w for w in winners if w['prize_id'] == p['id']]
        p['remaining'] = max(0, p['quantity'] - len(p['winners']))
    return jsonify({
        'prizes': prizes,
        'candidates': candidates,
        'has_data': bool(prizes and candidates)
    })

@app.route('/api/reset_winners', methods=['POST'])
@login_required
def reset_winners():
    conn = get_db()
    conn.execute("DELETE FROM winners")
    conn.commit()
    conn.close()
    return jsonify({'success': True})

@app.route('/uploads/<path:filename>')
def uploaded_file(filename):
    return send_from_directory(UPLOAD_DIR, filename)

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5670))
    app.run(host='0.0.0.0', port=port, debug=False)